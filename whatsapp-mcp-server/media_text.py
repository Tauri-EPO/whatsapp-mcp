"""Documents turned into text on this side of the wire, for ``read_media(as_text=True)``.

A clinical PDF, a contract in DOCX, a spreadsheet of readings: base64 gets them
to the agent, but 5 MiB of base64 is both over ``read_media``'s generic cap and
useless to a model that cannot decode a PDF. Extracting here turns the same file
into a few dozen KB of text.

One small library per format, chosen for being pure Python (``lxml`` apart) and
for doing nothing else: `pypdf` (page text), `python-docx` (paragraphs and
tables), `openpyxl` (cells, sheet by sheet). Deliberately out of scope:

* **no OCR.** A scanned PDF has no text layer, and this says so instead of
  returning an empty block that reads like an empty document;
* **no legacy .doc/.xls/.ppt** and no PPTX: different formats, different
  libraries, and the archive that motivated this carries PDFs and DOCX;
* **no layout**. Reading order per page, tabs between cells; a table is rows of
  text, not a rendering.

Everything is bounded: pages/sheets by ``max_pages``, rows per sheet by
:data:`MAX_ROWS_PER_SHEET`, and the whole answer by :data:`MAX_TEXT_CHARS`.
``truncated`` in the result says one of those cut the document short.
"""

from __future__ import annotations

import os
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field

from errors import ToolError

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXTRACTABLE_MIMES = frozenset({PDF_MIME, DOCX_MIME, XLSX_MIME})

# The file itself may be far bigger than what read_media returns as bytes (that
# is the point), but not unbounded: parsing a document is memory the server
# spends, and 64 MiB is already an outlier for a WhatsApp attachment.
MAX_EXTRACT_BYTES = 64 * 1024 * 1024
# DOCX and XLSX are zip containers and the parsers read their XML into memory,
# so the size on disk bounds nothing: 20 MiB of well-compressed zeroes expand
# into gigabytes. Anyone who can message the account chooses this file, so the
# declared uncompressed size of the members is checked before a parser sees it.
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
# What one call may put into the agent's context, whatever the document holds.
MAX_TEXT_CHARS = 200_000
MAX_PAGES = 20
MAX_PAGES_LIMIT = 500
MAX_ROWS_PER_SHEET = 500
# A cell holding a whole paragraph is a note, not a value.
MAX_CELL_CHARS = 500

TRUNCATION_NOTE = "[truncated: the rest of this document was not read]"
NO_TEXT_LAYER = (
    "This PDF has no extractable text: it is almost certainly a scan (photographed pages). "
    "There is no OCR on this server, so the words cannot be recovered here — "
    "read it as bytes with as_text=false, or ask the sender for the original file."
)


@dataclass
class _Cut:
    """Whether a row cap dropped part of a table or a sheet, set where it happens."""

    hit: bool = False


@dataclass
class Extracted:
    """The text of one document, plus what was left out."""

    sections: list[str] = field(default_factory=list)
    units_total: int = 0
    truncated: bool = False
    note: str | None = None


def page_limit(max_pages: int | None) -> int:
    """``max_pages`` as a positive count; 0 or unset means the default."""
    try:
        value = int(max_pages or 0)
    except (TypeError, ValueError) as exc:
        raise ToolError("invalid_argument", f"max_pages must be an integer, got {max_pages!r}") from exc
    if value <= 0:
        value = MAX_PAGES
    return min(value, MAX_PAGES_LIMIT)


def _guard_zip(path: str) -> None:
    """Refuse an office file that expands past what we will hold in memory.

    The check is the uncompressed size the central directory declares, which is
    enough because ``zipfile`` believes the same number: it stops the reader at
    ``file_size`` and raises ``BadZipFile`` on the CRC mismatch that follows
    (measured on CPython 3.13 with a member whose header was patched to lie:
    10 bytes came out of a 5 MiB payload, then BadZipFile). A lying header
    therefore cannot expand into memory through the parsers either, and the
    failure arrives as invalid_argument, not as a dead container.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except (OSError, zipfile.BadZipFile) as exc:
        raise ToolError("invalid_argument", f"this file is not a readable DOCX/XLSX container: {exc}") from exc
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ToolError(
            "too_large",
            f"this document expands to {total} bytes, more than as_text will parse ({MAX_UNCOMPRESSED_BYTES})",
            bytes=total,
            limit=MAX_UNCOMPRESSED_BYTES,
        )


def _guard_size(path: str) -> None:
    try:
        size = os.path.getsize(path)
    except OSError as exc:
        raise ToolError("internal", f"could not stat the document: {exc}") from exc
    if size > MAX_EXTRACT_BYTES:
        raise ToolError(
            "too_large",
            f"the document is {size} bytes and as_text parses at most {MAX_EXTRACT_BYTES}",
            bytes=size,
            limit=MAX_EXTRACT_BYTES,
        )


def _collect(sections: Iterator[str], produced: int, wanted: int) -> Extracted:
    """Take sections until the character budget runs out.

    ``produced`` is how many units (pages, sheets, blocks) the document holds
    and ``wanted`` how many this call asked for, so the caller learns that a
    20-page limit cut a 300-page file even when the text budget was never hit.
    """
    out = Extracted(units_total=produced, truncated=produced > wanted)
    used = 0
    for section in sections:
        if used + len(section) > MAX_TEXT_CHARS:
            room = MAX_TEXT_CHARS - used
            if room > len(TRUNCATION_NOTE):
                out.sections.append(section[: room - len(TRUNCATION_NOTE)] + TRUNCATION_NOTE)
            out.truncated = True
            break
        out.sections.append(section)
        used += len(section)
    return out


def _pdf(path: str, wanted: int) -> Extracted:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        pages = reader.pages
        total = len(pages)
        texts = [(index, (pages[index].extract_text() or "").strip()) for index in range(min(total, wanted))]
    except Exception as exc:  # noqa: BLE001 - a malformed file is bad input, not a server fault
        # pypdf raises PyPdfError subclasses for a broken file, but also
        # KeyError/ValueError from deep inside the parser on a truncated one.
        raise ToolError("invalid_argument", f"this PDF could not be read: {type(exc).__name__}: {exc}") from exc
    out = _collect((f"--- page {index + 1} of {total} ---\n{text}\n" for index, text in texts if text), total, wanted)
    if not out.sections:
        # Empty because there is nothing to read, or empty because the pages we
        # were allowed to read happen to be the image-only ones: telling a
        # 300-page report to go and ask the sender for the original would be wrong.
        out.note = (
            NO_TEXT_LAYER
            if total <= wanted
            else f"no text in the first {wanted} of {total} pages; raise max_pages to look further"
        )
    return out


def _docx(path: str, wanted: int) -> Extracted:
    import docx
    from docx.table import Table

    try:
        document = docx.Document(path)
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
        body = "\n".join(line for line in blocks if line)
        tables: list[Table] = list(document.tables)
    except Exception as exc:  # noqa: BLE001 - same as above: bad input, not a server fault
        raise ToolError("invalid_argument", f"this DOCX could not be read: {type(exc).__name__}: {exc}") from exc

    sections = [f"--- document text ---\n{body}\n"] if body else []
    total = len(sections) + len(tables)
    cut = _Cut()

    def rendered() -> Iterator[str]:
        # A generator, so max_pages=1 on a contract holding 200 tables does not
        # render 199 of them only to throw them away.
        yield from sections
        for number, table in enumerate(tables[: max(0, wanted - len(sections))], 1):
            yield f"--- table {number} of {len(tables)} ---\n" + _rows(_table_rows(table, cut))

    out = _collect(rendered(), total, wanted)
    out.truncated = out.truncated or cut.hit
    if not out.sections:
        out.note = "this DOCX has no text: it is empty, or holds only images"
    return out


def _table_rows(table, cut: _Cut) -> Iterator[list[str]]:
    rows = table.rows
    cut.hit = cut.hit or len(rows) > MAX_ROWS_PER_SHEET
    for row in rows[:MAX_ROWS_PER_SHEET]:
        yield [cell.text.replace("\n", " ").strip()[:MAX_CELL_CHARS] for cell in row.cells]


def _xlsx(path: str, wanted: int) -> Extracted:
    from openpyxl import load_workbook

    workbook = None
    cut = _Cut()
    rendered: list[str] = []
    total = 0
    try:
        # read_only defers the XML parsing to iter_rows, so a truncated file
        # fails inside the loop and not at load: both are bad input, not a
        # server fault, so the whole read is wrapped.
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = workbook.worksheets
        total = len(sheets)
        for sheet in sheets[:wanted]:
            cut.hit = cut.hit or (sheet.max_row or 0) > MAX_ROWS_PER_SHEET
            rows = _rows(
                [_cell(value) for value in row] for row in sheet.iter_rows(max_row=MAX_ROWS_PER_SHEET, values_only=True)
            )
            if rows:
                rendered.append(f"--- sheet {sheet.title} ({sheet.max_row or 0} rows) ---\n{rows}")
    except Exception as exc:  # noqa: BLE001 - same as above
        raise ToolError("invalid_argument", f"this XLSX could not be read: {type(exc).__name__}: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
    out = _collect(iter(rendered), total, wanted)
    out.truncated = out.truncated or cut.hit
    if not out.sections:
        out.note = "this spreadsheet has no cells with values"
    return out


def _cell(value: object) -> str:
    return "" if value is None else str(value).replace("\t", " ").replace("\n", " ")[:MAX_CELL_CHARS]


def _rows(rows) -> str:
    """Rows as tab-separated lines; empty rows are dropped, and nothing at all gives ""."""
    lines = [line for line in ("\t".join(cells) for cells in rows) if line.strip()]
    return "\n".join(lines) + "\n" if lines else ""


def require_extractable(mime: str) -> None:
    """Refuse a type ``as_text`` has no reader for, before anything is opened."""
    if mime not in EXTRACTABLE_MIMES:
        raise ToolError(
            "invalid_argument",
            f"as_text cannot read {mime}: it handles PDF, DOCX and XLSX (and text files are already text). "
            "Call again without as_text to get the bytes, or transcribe_audio for a voice note",
        )


def extract(path: str, mime: str, max_pages: int | None = None) -> Extracted:
    """The text of a PDF, DOCX or XLSX at ``path``; raises for anything else."""
    require_extractable(mime)
    _guard_size(path)
    wanted = page_limit(max_pages)
    if mime == PDF_MIME:
        return _pdf(path, wanted)
    _guard_zip(path)
    if mime == DOCX_MIME:
        return _docx(path, wanted)
    return _xlsx(path, wanted)
