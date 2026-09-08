"""read_media(as_text=True): PDF, DOCX and XLSX read on the server, as text.

The fixtures build real files with the three libraries (and a hand-written
minimal PDF, because pypdf writes pages but not words), so what is asserted is
what the extraction actually produces.
"""

import io
import json
import os

import pytest

import chat_policy
import main
import media_inventory
import media_read
import media_text
import whatsapp
from errors import ToolError
from tests.conftest import ALICE


def make_pdf(texts: list[str], pad: int = 0) -> bytes:
    """A minimal, valid PDF with one page per string in ``texts``."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * index for index in range(len(texts))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(texts)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for index, text in enumerate(texts):
        content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents {5 + 2 * index} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    start = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    if pad:
        # Bytes after %%EOF are ignored by the parser: a big file that still reads.
        out.write(b"%" + b"p" * pad + b"\n")
    return out.getvalue()


def make_docx(path: str) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("Contrato de prestação")
    document.add_paragraph("")
    document.add_paragraph("Cláusula primeira")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "item"
    table.cell(0, 1).text = "valor"
    table.cell(1, 0).text = "consulta"
    table.cell(1, 1).text = "250"
    document.save(path)


def make_xlsx(path: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    first = workbook.active
    first.title = "leituras"
    first.append(["data", "glicemia"])
    first.append(["2026-09-01", 98])
    second = workbook.create_sheet("notas")
    second.append(["jejum"])
    workbook.save(path)


def _insert(conn, msg_id, filename):
    conn.execute(
        "INSERT INTO messages (id, chat_jid, sender, content, timestamp, is_from_me, media_type, filename, "
        "file_length) VALUES (?, ?, '5511888888888', '', '2026-09-04 10:00:00+00:00', 0, 'document', ?, 0)",
        (msg_id, ALICE, filename),
    )


def _cache(name: str, data: bytes) -> str:
    directory = media_inventory.chat_media_dir(ALICE)
    os.makedirs(directory, exist_ok=True)
    path = os.path.join(directory, name)
    with open(path, "wb") as handle:
        handle.write(data)
    return path


@pytest.fixture
def documents(paired_dbs, monkeypatch):
    """One PDF, one scan-like PDF, one DOCX, one XLSX and one text file, all cached."""
    monkeypatch.setattr(whatsapp, "CHAT_POLICY", chat_policy.load_chat_policy({}))
    with paired_dbs.messages() as conn:
        for msg_id, filename in (
            ("PDF1", "laudo.pdf"),
            ("PDF2", "scan.pdf"),
            ("BIGPDF", "grande.pdf"),
            ("DOC1", "contrato.docx"),
            ("XLS1", "leituras.xlsx"),
            ("TXT1", "notas.txt"),
        ):
            _insert(conn, msg_id, filename)
    _cache("document_20260904_100000_PDF1.pdf", make_pdf(["Laudo do paciente", "Segunda pagina", "Terceira"]))
    _cache("document_20260904_100000_PDF2.pdf", make_pdf([""]))
    # Bigger than the 2 MiB base64 cap: as_text must read it anyway.
    _cache("document_20260904_100000_BIGPDF.pdf", make_pdf(["Exame completo"], pad=3 * 1024 * 1024))
    make_docx(_cache("document_20260904_100000_DOC1.docx", b""))
    make_xlsx(_cache("document_20260904_100000_XLS1.xlsx", b""))
    _cache("document_20260904_100000_TXT1.txt", b"plain notes\n")
    return paired_dbs


def _meta(blocks):
    return json.loads(blocks[-1].text)


def _texts(blocks):
    return [block.text for block in blocks[:-1]]


class TestPdf:
    def test_pages_come_back_as_text_with_markers(self, documents):
        blocks = media_read.read_media(ALICE, "PDF1", as_text=True)
        assert len(_texts(blocks)) == 3
        assert _texts(blocks)[0].startswith("--- page 1 of 3 ---")
        assert "Laudo do paciente" in _texts(blocks)[0]
        meta = _meta(blocks)
        assert meta["resource_link"]["uri"] == f"whatsapp://media/{ALICE}/PDF1"
        assert {key: value for key, value in meta.items() if key != "resource_link"} | {"notes": {}} == {
            "sha256": None,
            "mime": "application/pdf",
            "bytes": _meta(blocks)["bytes"],
            "truncated": False,
            "pages_total": 3,
            "notes": {},
        }

    def test_max_pages_stops_early_and_says_so(self, documents):
        blocks = media_read.read_media(ALICE, "PDF1", as_text=True, max_pages=1)
        assert len(_texts(blocks)) == 1
        assert _meta(blocks)["truncated"] is True and _meta(blocks)["pages_total"] == 3

    def test_a_scan_says_there_is_no_text_layer(self, documents):
        blocks = media_read.read_media(ALICE, "PDF2", as_text=True)
        assert "no extractable text" in _texts(blocks)[0]
        assert "OCR" in _texts(blocks)[0]

    def test_a_document_over_the_byte_cap_is_still_readable_as_text(self, documents):
        """The point of as_text: 3 MiB of PDF, a few hundred bytes of words."""
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "BIGPDF")
        assert exc.value.code == "too_large"
        blocks = media_read.read_media(ALICE, "BIGPDF", as_text=True)
        assert "Exame completo" in _texts(blocks)[0]
        # ...and no link to it: resources/read never extracts, so it would
        # answer too_large for the very file this call just read.
        assert "resource_link" not in _meta(blocks)

    def test_a_corrupt_pdf_is_the_caller_s_problem_not_an_internal_error(self, documents):
        _cache("document_20260904_100000_PDF1.pdf", b"%PDF-1.4 and then nothing useful")
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "PDF1", as_text=True)
        assert exc.value.code == "invalid_argument"

    def test_the_whole_answer_is_bounded(self, documents, monkeypatch):
        monkeypatch.setattr(media_text, "MAX_TEXT_CHARS", 40)
        blocks = media_read.read_media(ALICE, "PDF1", as_text=True)
        assert _meta(blocks)["truncated"] is True
        assert sum(len(text) for text in _texts(blocks)) <= 40


class TestBounds:
    """Every cap that drops content has to show up as truncated: an agent that
    sums a column must not be told the answer was complete when it was cut."""

    def test_a_sheet_cut_at_the_row_cap_says_truncated(self, documents, monkeypatch):
        monkeypatch.setattr(media_text, "MAX_ROWS_PER_SHEET", 1)
        blocks = media_read.read_media(ALICE, "XLS1", as_text=True)
        assert _meta(blocks)["truncated"] is True

    def test_a_table_cut_at_the_row_cap_says_truncated(self, documents, monkeypatch):
        monkeypatch.setattr(media_text, "MAX_ROWS_PER_SHEET", 1)
        blocks = media_read.read_media(ALICE, "DOC1", as_text=True)
        assert _meta(blocks)["truncated"] is True

    def test_max_pages_is_not_reported_as_a_missing_text_layer(self, documents):
        """The scan sentence tells the agent to give up; a page limit must not trigger it."""
        _cache("document_20260904_100000_PDF1.pdf", make_pdf(["", "", "texto na terceira"]))
        blocks = media_read.read_media(ALICE, "PDF1", as_text=True, max_pages=2)
        note = _texts(blocks)[0]
        assert "OCR" not in note and "max_pages" in note

    def test_an_empty_spreadsheet_says_so(self, documents):
        from openpyxl import Workbook

        path = _cache("document_20260904_100000_XLS1.xlsx", b"")
        Workbook().save(path)
        assert "no cells with values" in _texts(media_read.read_media(ALICE, "XLS1", as_text=True))[0]

    def test_a_zip_bomb_is_refused_before_a_parser_sees_it(self, documents, monkeypatch):
        monkeypatch.setattr(media_text, "MAX_UNCOMPRESSED_BYTES", 100)
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "DOC1", as_text=True)
        assert exc.value.code == "too_large"

    def test_a_header_that_lies_about_its_size_fails_instead_of_expanding(self, documents):
        """zipfile stops the reader at the declared size and then fails the CRC."""
        import struct
        import zipfile

        raw = io.BytesIO()
        with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"\0" * (4 * 1024 * 1024))
        data = bytearray(raw.getvalue())
        struct.pack_into("<I", data, data.rfind(b"PK\x01\x02") + 24, 10)
        struct.pack_into("<I", data, data.find(b"PK\x03\x04") + 22, 10)
        _cache("document_20260904_100000_DOC1.docx", bytes(data))
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "DOC1", as_text=True)
        assert exc.value.code == "invalid_argument"

    def test_a_corrupt_xlsx_is_invalid_argument_even_when_it_fails_while_reading(self, documents):
        """read_only parses the sheet inside iter_rows, long after load_workbook."""
        import zipfile

        path = _cache("document_20260904_100000_XLS1.xlsx", b"")
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "<nonsense/>")
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "XLS1", as_text=True)
        assert exc.value.code == "invalid_argument"


class TestOffice:
    def test_docx_gives_paragraphs_then_tables(self, documents):
        texts = _texts(media_read.read_media(ALICE, "DOC1", as_text=True))
        assert "Contrato de prestação" in texts[0] and "Cláusula primeira" in texts[0]
        assert texts[1].startswith("--- table 1 of 1 ---")
        assert "consulta\t250" in texts[1]

    def test_xlsx_gives_one_block_per_sheet_as_rows(self, documents):
        texts = _texts(media_read.read_media(ALICE, "XLS1", as_text=True))
        assert texts[0].startswith("--- sheet leituras")
        assert "data\tglicemia" in texts[0] and "2026-09-01\t98" in texts[0]
        assert texts[1].startswith("--- sheet notas")

    def test_max_pages_counts_sheets_too(self, documents):
        blocks = media_read.read_media(ALICE, "XLS1", as_text=True, max_pages=1)
        assert len(_texts(blocks)) == 1 and _meta(blocks)["truncated"] is True


class TestScope:
    def test_as_text_on_an_image_is_refused_rather_than_ignored(self):
        with pytest.raises(ToolError) as exc:
            media_text.extract("whatever", "image/jpeg")
        assert exc.value.code == "invalid_argument" and "PDF, DOCX and XLSX" in exc.value.message

    def test_as_text_on_a_big_video_says_as_text_not_too_large(self, documents):
        """The refusal has to name the real reason: no reader for video/mp4."""
        with documents.messages() as conn:
            _insert(conn, "VID1", "filme.mp4")
        _cache("document_20260904_100000_VID1.mp4", b"x" * (3 * 1024 * 1024))
        with pytest.raises(ToolError) as exc:
            media_read.read_media(ALICE, "VID1", as_text=True)
        assert exc.value.code == "invalid_argument" and "as_text cannot read video/mp4" in exc.value.message

    def test_as_text_on_a_text_file_changes_nothing(self, documents):
        blocks = media_read.read_media(ALICE, "TXT1", as_text=True)
        assert _texts(blocks) == ["plain notes\n"]
        assert "pages_total" not in _meta(blocks)

    def test_max_pages_is_validated_and_capped(self):
        assert media_text.page_limit(0) == media_text.MAX_PAGES
        assert media_text.page_limit(10_000) == media_text.MAX_PAGES_LIMIT
        with pytest.raises(ToolError):
            media_text.page_limit("many")  # type: ignore[arg-type]

    def test_the_tool_declares_both_new_arguments(self):
        """#282: an argument the tool does not declare is refused, so they must be declared."""
        from strict_args import declared_arguments

        tool = main.mcp._tool_manager.get_tool("read_media")
        assert tool is not None
        assert declared_arguments(tool) == {
            "chat_jid",
            "message_id",
            "max_bytes",
            "as_text",
            "max_pages",
            "max_edge",
            "quality",
        }
